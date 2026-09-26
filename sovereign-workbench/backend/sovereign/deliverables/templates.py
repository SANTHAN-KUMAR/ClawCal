"""Organisation templates: a new template is a file drop, not a code change.

Drop a Word file at `SOVEREIGN_DATA_DIR/templates/<kind>.docx` — `approval_note`
or `report` — and every deliverable of that kind is built on it: the
organisation's own styles, logo, letterhead, headers and footers. The template
marks where content goes with placeholders:

    {{org_name}} {{org_unit}} {{reference_no}} {{date}} {{subject}} {{title}}
    {{equipment_tag}} {{prepared_by}} {{audit_head}}      scalar fields
    {{body}}                                              where the note goes

Placeholders may sit in paragraphs, tables, headers or footers, and may be split
across formatting runs (Word does this when a word is edited twice); they are
matched on the paragraph's text, not on one run.

Check a template before relying on it:

    python -m sovereign.deliverables.templates validate approval_note.docx
    python -m sovereign.deliverables.templates init        # write starter templates
"""
from __future__ import annotations

import re
import sys
from pathlib import Path
from typing import Any, Iterator

from ..config import DATA_DIR

TEMPLATE_DIR = DATA_DIR / "templates"
KINDS = ("approval_note", "report")
KNOWN = {"org_name", "org_unit", "reference_no", "date", "subject", "title",
         "equipment_tag", "prepared_by", "audit_head", "body"}
REQUIRED = {"body"}
_PH = re.compile(r"\{\{\s*([a-z_][a-z0-9_]*)\s*\}\}")


def template_path(kind: str) -> Path | None:
    p = TEMPLATE_DIR / f"{kind}.docx"
    return p if p.is_file() else None


def _paragraphs(doc: Any) -> Iterator[Any]:
    """Every paragraph: body, tables (nested), headers and footers."""
    def walk(container: Any) -> Iterator[Any]:
        for p in container.paragraphs:
            yield p
        for t in getattr(container, "tables", []):
            for row in t.rows:
                for cell in row.cells:
                    yield from walk(cell)
    yield from walk(doc)
    for sec in doc.sections:
        for part in (sec.header, sec.footer, sec.first_page_header,
                     sec.first_page_footer, sec.even_page_header, sec.even_page_footer):
            if part is not None and not part.is_linked_to_previous:
                yield from walk(part)


def validate(path: str | Path) -> dict[str, Any]:
    from docx import Document
    try:
        doc = Document(str(path))
    except Exception as exc:
        return {"ok": False, "errors": [f"not a readable .docx: {exc}"],
                "placeholders": []}
    found: dict[str, int] = {}
    for p in _paragraphs(doc):
        for m in _PH.finditer(p.text):
            found[m.group(1)] = found.get(m.group(1), 0) + 1
    unknown = sorted(set(found) - KNOWN)
    missing = sorted(REQUIRED - set(found))
    warnings = []
    if "audit_head" not in found:
        warnings.append("no {{audit_head}}: the chain hash is still written to the "
                        "document properties and the generated footer, but not "
                        "where this template's reader will look")
    if found.get("body", 0) > 1:
        warnings.append("{{body}} appears more than once; content goes at the first")
    errors = []
    if missing:
        errors.append(f"missing required placeholder(s): "
                      f"{', '.join('{{' + m + '}}' for m in missing)} — the "
                      f"generated content would have nowhere to go")
    if unknown:
        errors.append(f"unknown placeholder(s) {unknown}; they would be printed "
                      f"literally. Known: {sorted(KNOWN)}")
    return {"ok": not errors, "placeholders": sorted(found), "counts": found,
            "errors": errors, "warnings": warnings}


def fill(doc: Any, values: dict[str, str]) -> None:
    """Replace scalar placeholders everywhere, keeping each paragraph's formatting."""
    for p in _paragraphs(doc):
        text = p.text
        if "{{" not in text or not _PH.search(text):
            continue
        new = _PH.sub(lambda m: str(values.get(m.group(1), m.group(0)))
                      if m.group(1) != "body" else m.group(0), text)
        if new == text:
            continue
        runs = p.runs
        if not runs:
            continue
        runs[0].text = new               # the first run's formatting carries it
        for r in runs[1:]:
            r.text = ""


def base_document(kind: str) -> tuple[Any, dict[str, Any]]:
    """(document, info) — the organisation's template if present, else blank."""
    from docx import Document
    p = template_path(kind)
    if p is None:
        return Document(), {"template": None}
    v = validate(p)
    if not v["ok"]:
        # A broken template must not silently produce a note with its content
        # missing; fall back to the built-in layout and say why.
        return Document(), {"template": str(p), "rejected": v["errors"]}
    return Document(str(p)), {"template": str(p), "placeholders": v["placeholders"]}


class BodyAnchor:
    """Collect content appended to the end of the document and move it to {{body}}."""

    def __init__(self, doc: Any) -> None:
        self.doc = doc
        self.body = doc.element.body
        self.marker = None
        for p in doc.paragraphs:
            if _PH.search(p.text) and "body" in {m.group(1) for m in _PH.finditer(p.text)}:
                self.marker = p._p
                break
        self.start = len(self.body)

    def place(self) -> None:
        if self.marker is None:
            return
        sect = self.body.find("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}sectPr")
        new = [el for el in list(self.body)[self.start - (1 if sect is not None else 0):]
               if el is not sect and el is not self.marker]
        for el in new:
            self.marker.addprevious(el)
        self.marker.getparent().remove(self.marker)


def starter(kind: str, out: Path) -> Path:
    """A starter template in the built-in look, for an organisation to edit in Word."""
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Pt
    doc = Document()
    for text, size, bold in (("{{org_name}}", 16, True), ("{{org_unit}}", 10.5, True),
                             ("CONFIDENTIAL — INTERNAL USE ONLY", 8.5, True)):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(text)
        r.bold, r.font.size = bold, Pt(size)
    t = doc.add_table(rows=2, cols=2)
    t.cell(0, 0).text, t.cell(0, 1).text = "Reference No.", "{{reference_no}}"
    t.cell(1, 0).text, t.cell(1, 1).text = "Date", "{{date}}"
    doc.add_paragraph()
    doc.add_paragraph("{{body}}")
    footer = doc.sections[0].footer.paragraphs[0]
    footer.text = "{{reference_no}} · audit chain head {{audit_head}}"
    out.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(out))
    return out


def main(argv: list[str]) -> int:
    import json
    if argv[:1] == ["validate"] and len(argv) == 2:
        r = validate(argv[1])
        print(json.dumps(r, indent=1))
        return 0 if r["ok"] else 1
    if argv[:1] == ["init"]:
        for k in KINDS:
            p = TEMPLATE_DIR / f"{k}.docx"
            if p.exists():
                print(f"kept {p}")
            else:
                print(f"wrote {starter(k, p)}")
        return 0
    print(__doc__)
    return 2


if __name__ == "__main__":
    sys.exit(main(sys.argv[1:]))
