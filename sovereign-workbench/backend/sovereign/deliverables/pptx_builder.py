"""PowerPoint deliverables: management and board briefing decks."""
from __future__ import annotations

import hashlib
import time
from pathlib import Path
from typing import Any

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.util import Inches, Pt

from .. import db
from ..config import ARTIFACT_DIR, settings

ACCENT = RGBColor(0x1F, 0x3A, 0x5F)
MUTED = RGBColor(0x55, 0x5F, 0x6B)
ALERT = RGBColor(0x99, 0x2B, 0x1F)


def _banner(slide: Any, prs: Presentation, text: str) -> None:
    box = slide.shapes.add_textbox(Inches(0.4), prs.slide_height - Inches(0.45),
                                   prs.slide_width - Inches(0.8), Inches(0.32))
    p = box.text_frame.paragraphs[0]
    r = p.add_run()
    r.text = text
    r.font.size = Pt(9)
    r.font.color.rgb = MUTED


def build_deck(title: str, slides: list[dict[str, Any]], *,
               subtitle: str = "", task_id: str | None = None,
               filename: str | None = None) -> dict[str, Any]:
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    # Title slide
    s = prs.slides.add_slide(prs.slide_layouts[6])
    tb = s.shapes.add_textbox(Inches(0.9), Inches(2.4), Inches(11.5), Inches(1.4))
    p = tb.text_frame.paragraphs[0]
    r = p.add_run()
    r.text = title
    r.font.size = Pt(40)
    r.font.bold = True
    r.font.color.rgb = ACCENT
    sb = s.shapes.add_textbox(Inches(0.95), Inches(3.8), Inches(11.5), Inches(1.4))
    for line in [subtitle or settings.org_unit, settings.org_name,
                 "CONFIDENTIAL — INTERNAL USE ONLY",
                 time.strftime("%d %B %Y")]:
        para = sb.text_frame.add_paragraph()
        run = para.add_run()
        run.text = line
        run.font.size = Pt(14)
        run.font.color.rgb = ALERT if "CONFIDENTIAL" in line else MUTED
        if "CONFIDENTIAL" in line:
            run.font.bold = True

    for idx, spec in enumerate(slides, 1):
        s = prs.slides.add_slide(prs.slide_layouts[6])
        hb = s.shapes.add_textbox(Inches(0.6), Inches(0.4), Inches(12.1), Inches(0.9))
        hp = hb.text_frame.paragraphs[0]
        hr = hp.add_run()
        hr.text = spec.get("heading", f"Slide {idx}")
        hr.font.size = Pt(26)
        hr.font.bold = True
        hr.font.color.rgb = ACCENT

        bullets = spec.get("bullets") or []
        if bullets:
            bb = s.shapes.add_textbox(Inches(0.8), Inches(1.5), Inches(11.7),
                                      Inches(5.0))
            tf = bb.text_frame
            tf.word_wrap = True
            for i, b in enumerate(bullets):
                para = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
                run = para.add_run()
                text = b if isinstance(b, str) else str(b)
                run.text = f"•  {text}"
                run.font.size = Pt(16)
                para.space_after = Pt(10)
                if text.startswith("CANNOT DETERMINE") or "[CANNOT DETERMINE" in text:
                    run.font.color.rgb = ALERT
                    run.font.bold = True

        table = spec.get("table")
        if table:
            cols = spec.get("headers") or sorted({k for r_ in table for k in r_})
            top = Inches(1.5 + (0.35 * len(bullets) if bullets else 0))
            shape = s.shapes.add_table(len(table) + 1, len(cols), Inches(0.8), top,
                                       Inches(11.7),
                                       Inches(0.4 * (len(table) + 1)))
            tbl = shape.table
            for c, name in enumerate(cols):
                cell = tbl.cell(0, c)
                cell.text = str(name)
                cell.text_frame.paragraphs[0].runs[0].font.size = Pt(12)
                cell.text_frame.paragraphs[0].runs[0].font.bold = True
            for r_i, row in enumerate(table, 1):
                for c, name in enumerate(cols):
                    cell = tbl.cell(r_i, c)
                    cell.text = str(row.get(name, ""))
                    for para in cell.text_frame.paragraphs:
                        for run in para.runs:
                            run.font.size = Pt(11)

        note = spec.get("note")
        if note:
            nb = s.shapes.add_textbox(Inches(0.8), Inches(6.3), Inches(11.7),
                                      Inches(0.7))
            np_ = nb.text_frame.paragraphs[0]
            nr = np_.add_run()
            nr.text = note
            nr.font.size = Pt(11)
            nr.font.italic = True
            nr.font.color.rgb = MUTED

        _banner(s, prs, f"{settings.org_name}  |  CONFIDENTIAL  |  slide {idx}")

    out = ARTIFACT_DIR / (filename or f"deck_{int(time.time())}.pptx")
    prs.save(str(out))
    data = out.read_bytes()
    aid = db.new_id("art")
    db.insert("artifacts", {
        "id": aid, "task_id": task_id, "name": out.name, "kind": "presentation",
        "path": str(out), "bytes": len(data),
        "sha256": hashlib.sha256(data).hexdigest(),
        "meta": db.jdump({"title": title, "slides": len(slides)}),
        "created_at": time.time()})
    return {"artifact_id": aid, "name": out.name, "path": str(out),
            "bytes": len(data), "kind": "presentation", "slides": len(slides)}
