# -*- coding: utf-8 -*-
"""Rebuild the SIH 'Technical Approach' slide as native, editable PowerPoint shapes."""
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR, MSO_AUTO_SIZE
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.enum.dml import MSO_LINE_DASH_STYLE
from pptx.oxml.ns import qn
from lxml import etree

# ---------- palette (sampled from the existing deck) ----------
INK    = RGBColor(0x1A, 0x1A, 0x1A)
MUTE   = RGBColor(0x6B, 0x70, 0x76)
BODY   = RGBColor(0x4A, 0x4F, 0x55)
RULE   = RGBColor(0xC4, 0xC8, 0xCC)
HAIR   = RGBColor(0x9A, 0xA1, 0xA8)
FILL   = RGBColor(0xF1, 0xF2, 0xF3)
FILL2  = RGBColor(0xF5, 0xF6, 0xF7)
EDGE   = RGBColor(0xDD, 0xE0, 0xE3)
RED    = RGBColor(0xC0, 0x00, 0x00)   # deck heading red
TILE   = RGBColor(0xA8, 0x30, 0x2B)   # deck "we build it" tile red
SLATE  = RGBColor(0x5C, 0x6A, 0x75)   # deck "assemble" tile
BLUE   = RGBColor(0x1F, 0x6F, 0xB5)   # deck footer bar
PURPLE = RGBColor(0x7C, 0x5F, 0xA6)
WHITE  = RGBColor(0xFF, 0xFF, 0xFF)
ONRED  = RGBColor(0xF0, 0xD5, 0xD3)   # muted text on the red tiles
REDLN  = RGBColor(0xCE, 0x7B, 0x76)

SANS, SERIF, MONO = "Arial", "Times New Roman", "Consolas"

prs = Presentation()
prs.slide_width, prs.slide_height = Inches(13.3333), Inches(7.5)
slide = prs.slides.add_slide(prs.slide_layouts[6])
S = slide.shapes

def E(px):            # design px (1280x720 canvas) -> EMU at 96 px/in
    return Emu(int(round(px / 96 * 914400)))

def _noline(sh):
    sh.line.fill.background()

def rect(x, y, w, h, fill=None, line=None, lw=1.0, dash=False, shape=MSO_SHAPE.RECTANGLE):
    sh = S.add_shape(shape, E(x), E(y), E(w), E(h))
    sh.shadow.inherit = False
    if fill is None:
        sh.fill.background()
    else:
        sh.fill.solid(); sh.fill.fore_color.rgb = fill
    if line is None:
        _noline(sh)
    else:
        sh.line.color.rgb = line; sh.line.width = Pt(lw)
        if dash:
            sh.line.dash_style = MSO_LINE_DASH_STYLE.DASH
    sh.text_frame.text = ""
    return sh

def text(x, y, w, h, runs, size=9, color=INK, bold=False, font=SANS,
         align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, spc=None, line_spc=None, caps=False):
    """runs: a string, or a list of (text, overrides-dict)."""
    tb = S.add_textbox(E(x), E(y), E(w), E(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.auto_size = MSO_AUTO_SIZE.NONE
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = anchor
    if isinstance(runs, str):
        runs = [(runs, {})]
    p = tf.paragraphs[0]
    p.alignment = align
    if line_spc:
        p.line_spacing = line_spc
    for t, ov in runs:
        r = p.add_run()
        r.text = t.upper() if ov.get("caps", caps) else t
        f = r.font
        f.name  = ov.get("font", font)
        f.size  = Pt(ov.get("size", size))
        f.bold  = ov.get("bold", bold)
        f.color.rgb = ov.get("color", color)
        sp = ov.get("spc", spc)
        if sp:
            r._r.get_or_add_rPr().set("spc", str(int(sp * 100)))
        # make the East-Asian/complex-script fonts match so PowerPoint doesn't substitute
        rPr = r._r.get_or_add_rPr()
        for tag in ("a:ea", "a:cs"):
            el = etree.SubElement(rPr, qn(tag)); el.set("typeface", f.name)
    return tb

def line(x1, y1, x2, y2, color=HAIR, lw=1.0, arrow=False, dash=False):
    c = S.add_connector(MSO_CONNECTOR.STRAIGHT, E(x1), E(y1), E(x2), E(y2))
    c.line.color.rgb = color
    c.line.width = Pt(lw)
    if dash:
        c.line.dash_style = MSO_LINE_DASH_STYLE.DASH
    if arrow:
        ln = c.line._get_or_add_ln()
        te = etree.SubElement(ln, qn("a:tailEnd"))
        te.set("type", "triangle"); te.set("w", "med"); te.set("len", "med")
    return c

def icon(name, variant, x, y, sz=16):
    return S.add_picture("icons/%s-%s.png" % (name, variant), E(x), E(y), E(sz), E(sz))

# =====================================================================
# HEADER
# =====================================================================
ov = rect(44, 20, 118, 52, fill=None, line=PURPLE, lw=1.25, shape=MSO_SHAPE.OVAL)
text(44, 20, 118, 52, "Dilijens", size=11, color=RGBColor(0x3F,0x33,0x50),
     align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)

text(0, 24, 1280, 46, "TECHNICAL APPROACH", size=30, bold=True, font=SERIF,
     align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE, spc=1.2)

rect(1084, 18, 152, 56, fill=None, line=EDGE, lw=0.75, dash=True)
text(1084, 18, 152, 56, "paste SIH 2026 logo", size=7, color=RGBColor(0xB9,0xBE,0xC3),
     font=MONO, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)

text(44, 82, 800, 16, "WHAT SITS ON THE BOX, AND WHY IT FITS ON ONE",
     size=8.5, bold=True, color=RED, spc=1.0)

# =====================================================================
# BAND 1 — ARCHITECTURE   (content band y 134..394)
# =====================================================================
text(44, 102, 200, 18, "The plant network", size=12, bold=True)
line(44, 125, 212, 125, INK, 1.25)

text(246, 102, 620, 18, "The workbench — one on-premise server", size=12, bold=True)
line(246, 125, 1058, 125, INK, 1.25)
rect(928, 109, 8, 8, fill=TILE)
text(942, 103, 130, 16, "red is what we write", size=7.5, color=MUTE, anchor=MSO_ANCHOR.MIDDLE)

# ---- clients -------------------------------------------------------
text(44, 176, 168, 12, "MANY CLIENTS, ONE SERVER", size=7, bold=True, color=MUTE, spc=0.9)
for i, (ic, label) in enumerate([("monitor", "engineer’s desk"),
                                 ("tablet",  "inspector’s tablet"),
                                 ("clock",   "a scheduled job"),
                                 ("server",  "another system")]):
    cy = 192 + i * 37
    rect(44, cy, 168, 30, fill=FILL, line=EDGE, lw=0.75)
    icon(ic, "dark", 54, cy + 7)
    text(78, cy, 128, 30, label, size=9, anchor=MSO_ANCHOR.MIDDLE)
text(44, 338, 168, 14, "browser · CLI · REST", size=7.5, font=MONO, color=MUTE)

line(215, 264, 243, 264, HAIR, 1.1, arrow=True)

# ---- the agent -----------------------------------------------------
rect(246, 134, 232, 144, fill=WHITE, line=RULE, lw=0.75)
text(258, 145, 208, 16, "AGENT LOOP", size=10.5, bold=True, spc=0.3)
px = 258
for w, lbl in ((38, "plan"), (34, "act"), (58, "observe")):
    rect(px, 172, w, 20, fill=None, line=RULE, lw=0.75)
    text(px, 172, w, 20, lbl, size=9, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    px += w
    if lbl != "observe":
        text(px, 172, 8, 20, "›", size=11, color=MUTE, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
        px += 8
line(396, 206, 262, 206, HAIR, 1.0, arrow=True)
rect(300, 198, 58, 16, fill=WHITE)
text(300, 198, 58, 16, "repeat", size=9, color=MUTE, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
text(258, 240, 208, 14, "goose · MCP tools", size=7.5, font=MONO, color=MUTE)

rect(246, 292, 232, 102, fill=WHITE, line=RULE, lw=0.75)
text(258, 302, 208, 16, "SCHEDULER & STATE", size=10.5, bold=True, spc=0.3)
text(258, 324, 208, 18, "cron · long runs · replayable history", size=9)
text(258, 356, 208, 14, "Temporal · Postgres", size=7.5, font=MONO, color=MUTE)

# fork: agent -> the two checks
line(478, 206, 504, 206, HAIR, 1.1, arrow=True)
line(486, 206, 486, 343, HAIR, 1.1)
line(486, 343, 504, 343, HAIR, 1.1, arrow=True)

# ---- the checks (ours) ---------------------------------------------
rect(510, 134, 208, 144, fill=TILE)
text(522, 145, 184, 16, "APPROVAL GATE", size=10.5, bold=True, color=WHITE, spc=0.3)
text(522, 164, 184, 14, "+ provenance", size=8.5, color=ONRED)
line(522, 187, 706, 187, REDLN, 0.75)
text(522, 197, 184, 34, "Nothing acts or cites unchecked.", size=9.5, color=WHITE, line_spc=1.25)
text(522, 248, 184, 14, "we write it", size=7.5, font=MONO, color=ONRED)

rect(510, 292, 208, 102, fill=TILE)
text(522, 301, 184, 16, "ADMISSION ROUTER", size=10.5, bold=True, color=WHITE, spc=0.3)
text(522, 322, 184, 32, "Residency-aware, over the LiteLLM gateway.", size=9, color=WHITE, line_spc=1.25)
text(522, 368, 184, 14, "we write it", size=7.5, font=MONO, color=ONRED)

line(718, 206, 746, 206, HAIR, 1.1, arrow=True)
line(718, 343, 746, 343, HAIR, 1.1, arrow=True)

# ---- tool surface + model plane ------------------------------------
text(750, 134, 308, 12, "TOOL SURFACE", size=7, bold=True, color=MUTE, spc=0.9)
tools = [("terminal", "code sandbox",     TILE,  750, 152),
         ("globe",    "headless browser", SLATE, 908, 152),
         ("book",     "plant knowledge",  SLATE, 750, 216),
         ("doc",      "documents out",    SLATE, 908, 216)]
for ic, label, col, tx, ty in tools:
    rect(tx, ty, 150, 56, fill=col)
    icon(ic, "white", tx + 11, ty + 20)
    text(tx + 34, ty, 106, 56, label, size=9, color=WHITE, anchor=MSO_ANCHOR.MIDDLE)

text(750, 292, 308, 12, "MODEL PLANE", size=7, bold=True, color=MUTE, spc=0.9)
for nm, desc, mx in (("FreeToken", "MoE generalist", 750), ("vLLM", "small specialists", 908)):
    rect(mx, 310, 150, 82, fill=WHITE, line=RULE, lw=0.75)
    text(mx + 10, 319, 130, 14, nm, size=10, bold=True)
    text(mx + 10, 337, 130, 14, desc, size=8.5, color=BODY)
    rect(mx + 10, 360, 62, 16, fill=FILL)
    text(mx + 10, 360, 62, 16, "Apache-2.0", size=7, font=MONO, color=BODY,
         align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)

# ---- the fence -----------------------------------------------------
line(1058, 206, 1084, 206, HAIR, 1.1, arrow=True)
line(1090, 200, 1100, 212, RED, 2.0)
line(1100, 200, 1090, 212, RED, 2.0)

rect(1104, 134, 5, 260, fill=TILE)
text(1120, 206, 116, 16, "THE FENCE", size=10, bold=True, color=TILE, spc=0.7)
text(1120, 228, 116, 14, "default-DROP", size=7.5, font=MONO, color=BODY)
text(1120, 242, 116, 14, "nftables · Tetragon", size=7.5, font=MONO, color=BODY)
text(1120, 262, 122, 62, "Everything left of this line stays there. "
     "Nothing dials out — the kernel refuses.", size=9, line_spc=1.3)

# =====================================================================
# BAND 2 — THE INVERSION
# =====================================================================
text(44, 402, 800, 18, "The inversion — nothing swaps, on either machine", size=12, bold=True)
line(44, 425, 1236, 425, INK, 1.25)

def minibox(x, y, w, h, label, col=HAIR, txt=INK):
    rect(x, y, w, h, fill=None, line=col, lw=0.75)
    text(x, y, w, h, label, size=7, font=MONO, color=txt,
         align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)

# 01 — the obvious build (rejected)
rect(44, 434, 386, 104, fill=FILL2, line=EDGE, lw=0.75)
text(56, 445, 22, 14, "01", size=8.5, font=MONO, color=HAIR)
text(80, 445, 200, 14, "THE OBVIOUS BUILD", size=7, bold=True, color=MUTE, spc=1.0)
minibox(56, 466, 44, 18, "GPU", RULE, MUTE)
text(106, 466, 200, 18, "model A  ⇄  model B", size=7.5, font=MONO, color=MUTE, anchor=MSO_ANCHOR.MIDDLE)
text(56, 494, 362, 34, "One model fits. The swap costs 30–100 s, every agent step.",
     size=9, color=MUTE, line_spc=1.3)

# 02 — discrete card
rect(447, 434, 386, 104, fill=WHITE, line=RULE, lw=0.75)
text(459, 445, 22, 14, "02", size=8.5, font=MONO, color=HAIR)
text(483, 445, 220, 14, "OURS · DISCRETE CARD", size=7, bold=True, color=RED, spc=1.0)
minibox(459, 466, 68, 18, "HOST RAM")
text(531, 466, 56, 18, "→ PCIe →", size=7.5, font=MONO, color=MUTE,
     align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
minibox(591, 466, 96, 18, "GPU · hot + KV")
text(459, 494, 362, 34, "Experts stream across the bus. The transfer is the cost.",
     size=9, line_spc=1.3)

# 03 — unified memory
rect(850, 434, 386, 104, fill=WHITE, line=RULE, lw=0.75)
text(862, 445, 22, 14, "03", size=8.5, font=MONO, color=HAIR)
text(886, 445, 220, 14, "OURS · UNIFIED MEMORY", size=7, bold=True, color=RED, spc=1.0)
minibox(862, 466, 362, 18, "ONE POOL  ·  CPU + GPU  ·  KV")
text(862, 494, 362, 34, "No PCIe hop. Nothing to transfer — the pool is the ceiling.",
     size=9, line_spc=1.3)

# =====================================================================
# BAND 3 — THE CLAIM
# =====================================================================
rect(44, 548, 1192, 46, fill=None, line=INK, lw=1.25)
text(44, 548, 1192, 46,
     "The scarce resource stops being VRAM. It becomes memory the GPU can reach, and how fast.",
     size=12.5, bold=True, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
text(44, 600, 1192, 14, "First run measures which of the two this machine is.",
     size=8, color=MUTE, align=PP_ALIGN.RIGHT)

# =====================================================================
# FOOTER (deck template)
# =====================================================================
rect(0, 648, 1280, 38, fill=BLUE)
text(0, 648, 1280, 38, "@SIH Idea submission- Template", size=9, color=WHITE,
     align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
text(1180, 648, 56, 38, "3", size=10, bold=True, color=WHITE,
     align=PP_ALIGN.RIGHT, anchor=MSO_ANCHOR.MIDDLE)

prs.save("SIH-Technical-Approach-slide3.pptx")
print("saved SIH-Technical-Approach-slide3.pptx ; shapes =", len(S))
