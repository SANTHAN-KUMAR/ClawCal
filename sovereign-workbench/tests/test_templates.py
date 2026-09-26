"""Organisation templates: a file drop, validated, with the audit head embedded."""
from __future__ import annotations

import time

from docx import Document

from sovereign.deliverables import docx_builder, templates


def _texts(path) -> str:
    d = Document(str(path))
    parts = [p.text for p in d.paragraphs]
    for t in d.tables:
        parts += [c.text for r in t.rows for c in r.cells]
    for s in d.sections:
        parts += [p.text for p in s.footer.paragraphs]
    return "\n".join(parts)


def test_validator_reports_placeholders_and_refuses_a_template_without_body(tmp_path):
    good = templates.starter("approval_note", tmp_path / "good.docx")
    v = templates.validate(good)
    assert v["ok"] and {"body", "org_name", "reference_no", "audit_head"} <= set(v["placeholders"])
    bad = tmp_path / "bad.docx"
    d = Document(); d.add_paragraph("{{org_name}} {{nonsense}}"); d.save(str(bad))
    v = templates.validate(bad)
    assert not v["ok"] and any("body" in e for e in v["errors"]) \
        and any("nonsense" in e for e in v["errors"])


def test_a_dropped_template_is_used_and_filled(tmp_path, monkeypatch):
    tdir = tmp_path / "templates"
    monkeypatch.setattr(templates, "TEMPLATE_DIR", tdir)
    t = templates.starter("approval_note", tdir / "approval_note.docx")
    # Word splits a placeholder across runs when it is edited; build one that way.
    d = Document(str(t))
    p = d.add_paragraph()
    p.add_run("Prepared for {{sub"); p.add_run("ject}} by the unit")
    d.save(str(t))
    art = docx_builder.build_approval_note(
        docx_builder.ApprovalNoteData(subject="V-204 continued operation",
                                      equipment_tag="V-204"),
        filename=f"tpl-{time.time_ns()}.docx")
    text = _texts(art["path"])
    assert art["template"] and "{{" not in text
    assert "V-204 continued operation" in text and "audit chain head" in text
    # The generated note sits where {{body}} was: after the template's
    # reference table, not appended below everything the template holds.
    body = list(Document(art["path"]).element.body)
    tbl = next(i for i, el in enumerate(body) if el.tag.endswith("}tbl"))
    head = next(i for i, el in enumerate(body)
                if "EQUIPMENT AND REFERENCE" in "".join(el.itertext()).upper())
    assert tbl < head


def test_without_a_template_the_built_in_layout_embeds_the_audit_head(tmp_path, monkeypatch):
    monkeypatch.setattr(templates, "TEMPLATE_DIR", tmp_path / "none")
    art = docx_builder.build_report("Weekly summary", [{"heading": "One", "body": "x"}],
                                    filename=f"r-{time.time_ns()}.docx")
    d = Document(art["path"])
    assert d.core_properties.keywords.startswith("clawcal audit-head ")
    assert "audit chain head" in _texts(art["path"])
